"""Write the M4 verification oracle for the reservoir curve (B3.3; OI-01, OI-05).

Run `python tests/verification/make_reservoir_oracle.py` to rewrite these files next to this script:

- `reservoir_verification.xlsx`: the hand calculation as live spreadsheet formulas, for the hydrologist to open and
  review. A case sheet holds the fit (dead volume, exponent m, Hmin, surface area at Hmax) and a point sheet holds the
  curve at five levels of each case. The shape-class exponents sit in their own cells.
- `reservoir_cases.csv`, `reservoir_points.csv`: the same cases with the results worked out below using only `math`.
  The unit tests read these files, so the oracle never depends on `src`.

The first three cases are the worked examples of B3.3. Change a case or a constant here, review the workbook, and only
then commit the files.
"""

import csv
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

SHAPE_CLASSES = {"gorge": 1.25, "hill": 2.0, "foothill": 3.0, "lake": 4.0}

# (case, description, Hbed m, Hmin m or None, Hmax m, Hcrest m, Vactive hm3, Vmax hm3, shape class or None)
CASES = [
    (1, "Kavsakkaya sample, foothill", 1063.0, None, 1122.3, 1124.0, 74.1, 80.8, "foothill"),
    (2, "Cubuk 1 sample, hill", 888.0, None, 907.6, 908.6, 5.28, 5.6, "hill"),
    (3, "Cubuk 2 sample, hill (OI-01)", 1110.0, None, 1115.3, 1117.0, 22.0, 22.4, "hill"),
    (4, "Kavsakkaya sample, Hmin given", 1063.0, 1090.0, 1122.3, 1124.0, 74.1, 80.8, None),
    (5, "Cubuk 1 sample, Hmin given", 888.0, 895.0, 907.6, 908.6, 5.28, 5.6, None),
    (6, "Round numbers, Hmin gives m = 2", 100.0, 125.0, 150.0, 155.0, 75.0, 100.0, None),
    (7, "Round numbers, gorge", 0.0, None, 10.0, 12.0, 90.0, 100.0, "gorge"),
    (8, "Round numbers, lake", 0.0, None, 10.0, 12.0, 90.0, 100.0, "lake"),
]

POINTS = ["Hbed", "Hmin", "Hmid", "Hmax", "Hcrest"]


def fitted(case):
    """Dead volume, exponent m, the Hmin of the curve and the surface area at Hmax (hm3 per m is 10^6 m2, so km2)."""
    _, _, bed, h_min, h_max, _, active, v_max, shape = case
    dead = v_max - active
    if h_min is None:
        m = SHAPE_CLASSES[shape]
        h_min = bed + (h_max - bed) * (dead / v_max) ** (1.0 / m)
    else:
        m = math.log(dead / v_max) / math.log((h_min - bed) / (h_max - bed))
    return dead, m, h_min, m * v_max / (h_max - bed)


def point_levels(case, h_min):
    _, _, bed, _, h_max, h_crest, _, _, _ = case
    return [bed, h_min, (bed + h_max) / 2.0, h_max, h_crest]


def volume(case, m, h):
    _, _, bed, _, h_max, _, _, v_max, _ = case
    return v_max * ((h - bed) / (h_max - bed)) ** m


def area(case, m, h):
    """dV/dH = m V / (H - Hbed), written as m Vmax / (Hmax - Hbed) x^(m-1) so that the bed itself is defined."""
    _, _, bed, _, h_max, _, _, v_max, _ = case
    return m * v_max / (h_max - bed) * ((h - bed) / (h_max - bed)) ** (m - 1.0)


def write_csv():
    with open(HERE / "reservoir_cases.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Case",
                "Description",
                "Hbed_m",
                "Hmin_given_m",
                "Hmax_m",
                "Hcrest_m",
                "Vactive_hm3",
                "Vmax_hm3",
                "Shape_Class",
                "Vdead_hm3",
                "m",
                "Hmin_m",
                "Area_Hmax_km2",
            ]
        )
        for case in CASES:
            number, description, bed, h_min, h_max, h_crest, active, v_max, shape = case
            dead, m, h_min_used, area_max = fitted(case)
            writer.writerow(
                [number, description, bed, "" if h_min is None else h_min, h_max, h_crest, active, v_max]
                + ["" if shape is None else shape, dead, m, h_min_used, area_max]
            )

    with open(HERE / "reservoir_points.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Case", "Point", "H_m", "V_hm3", "A_km2"])
        for case in CASES:
            _, m, h_min_used, _ = fitted(case)
            for name, h in zip(POINTS, point_levels(case, h_min_used), strict=True):
                writer.writerow([case[0], name, h, volume(case, m, h), area(case, m, h)])


def write_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "B3.3 cases"
    headers = [
        "Case",
        "Description",
        "Hbed (m)",
        "Hmin given (m, blank = shape class)",
        "Hmax (m)",
        "Hcrest (m)",
        "Vactive (hm3)",
        "Vmax (hm3)",
        "Shape class",
        "Shape-class exponent",
        "Vdead (hm3) = Vmax - Vactive",
        "m = ln(Vdead/Vmax) / ln((Hmin-Hbed)/(Hmax-Hbed)), or the class exponent",
        "Hmin used (m) = Hbed + (Hmax-Hbed) (Vdead/Vmax)^(1/m)",
        "A(Hmax) (km2) = m Vmax / (Hmax-Hbed)",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row, (number, description, bed, h_min, h_max, h_crest, active, v_max, shape) in enumerate(CASES, start=2):
        sheet.append(
            [
                number,
                description,
                bed,
                h_min,
                h_max,
                h_crest,
                active,
                v_max,
                shape,
                f'=IF(I{row}="","",VLOOKUP(I{row},$Q$2:$R$5,2,FALSE))',
                f"=H{row}-G{row}",
                f'=IF(D{row}="",J{row},LN(K{row}/H{row})/LN((D{row}-C{row})/(E{row}-C{row})))',
                f"=C{row}+(E{row}-C{row})*(K{row}/H{row})^(1/L{row})",
                f"=L{row}*H{row}/(E{row}-C{row})",
            ]
        )
    sheet["Q1"], sheet["R1"] = "Shape class", "Exponent m"
    for row, (name, exponent) in enumerate(SHAPE_CLASSES.items(), start=2):
        sheet.cell(row=row, column=17, value=name)
        sheet.cell(row=row, column=18, value=exponent)
    for column, width in zip("ABCDEFGHIJKLMN", (6, 32, 10, 20, 10, 10, 12, 12, 12, 14, 18, 36, 34, 26), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["Q"].width = 14

    points = book.create_sheet("B3.3 points")
    points.append(
        [
            "Case",
            "Point",
            "H (m)",
            "V (hm3) = Vmax x^m, x = (H-Hbed)/(Hmax-Hbed)",
            "A (km2) = m Vmax / (Hmax-Hbed) x^(m-1)",
            "H back from V (m) = Hbed + (Hmax-Hbed) (V/Vmax)^(1/m)",
        ]
    )
    for cell in points[1]:
        cell.font = Font(bold=True)
    row = 2
    for case_row, case in enumerate(CASES, start=2):

        def c(column, case_row=case_row):
            """A cell of this case on the case sheet."""
            return f"'B3.3 cases'!${column}${case_row}"

        levels = {
            "Hbed": f"={c('C')}",
            "Hmin": f"={c('M')}",
            "Hmid": f"=({c('C')}+{c('E')})/2",
            "Hmax": f"={c('E')}",
            "Hcrest": f"={c('F')}",
        }
        for name in POINTS:
            x = f"((C{row}-{c('C')})/({c('E')}-{c('C')}))"
            points.append(
                [
                    case[0],
                    name,
                    levels[name],
                    f"={c('H')}*{x}^{c('L')}",
                    f"={c('L')}*{c('H')}/({c('E')}-{c('C')})*{x}^({c('L')}-1)",
                    f"={c('C')}+({c('E')}-{c('C')})*(D{row}/{c('H')})^(1/{c('L')})",
                ]
            )
            row += 1
    for column, width in zip("ABCDEF", (6, 10, 12, 40, 36, 42), strict=True):
        points.column_dimensions[column].width = width

    book.calculation.fullCalcOnLoad = True  # openpyxl stores no results, so Excel works them out on opening
    book.save(HERE / "reservoir_verification.xlsx")


if __name__ == "__main__":
    write_csv()
    write_workbook()
