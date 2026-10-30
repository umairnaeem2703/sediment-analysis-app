"""Write the M4.2 verification oracle for the integration, accumulation and lifespan of B3.4 (OI-03).

Run `python tests/verification/make_integration_oracle.py` to rewrite these files next to this script:

- `integration_verification.xlsx`: the hand calculation as live spreadsheet formulas, for the hydrologist to open and
  review. Each case has its own sheet: the inputs at the top, then the annual totals and elevations of the record,
  the three rates, the accumulation to date year by year, the lifespans and the projection at five horizons.
- `integration_cases.csv`, `integration_record.csv`, `integration_accumulation.csv`, `integration_projection.csv`: the
  same cases with the results worked out below using only `math`. The unit tests read these files, so the oracle never
  depends on `src`.

All seven cases share one reservoir (Hbed 100 m, Hmax 150 m, Hcrest 155 m, Vactive 120 hm3, Vmax 200 hm3, m = 2, so
Vdead = 80 hm3 and V(Hcrest) = 242 hm3) and one five-year record, 2001 to 2005. Change a case or a constant here,
review the workbook, and only then commit the files.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

M3_PER_HM3 = 1_000_000.0
HBED, HMAX, HCREST, VACTIVE, VMAX, M = 100.0, 150.0, 155.0, 120.0, 200.0, 2.0
VDEAD = VMAX - VACTIVE
VCREST = VMAX * ((HCREST - HBED) / (HMAX - HBED)) ** M
HMIN = HBED + (HMAX - HBED) * (VDEAD / VMAX) ** (1.0 / M)

RECORD_YEARS = [2001, 2002, 2003, 2004, 2005]
BEDLOAD = [1_000_000.0, 2_500_000.0, 500_000.0, 2_000_000.0, 1_500_000.0]  # m3/yr
SPARSE = [0.0, 1_000_000.0, 2_000_000.0, 500_000.0, 1_500_000.0]  # the first year delivers nothing
RATES = ["Mean", "Min", "Max"]
HORIZONS = [0, 1, 10, 50, 100]  # years after the analysis year

# (case, description, suspended m3/yr, bedload m3/yr per record year, TE %, commissioning year, analysis year)
CASES = [
    (1, "Record years only, no trap loss", 4_000_000.0, BEDLOAD, 100.0, 2001, 2006),
    (2, "TE 60 percent, years outside the record at the mean", 4_000_000.0, BEDLOAD, 60.0, 1999, 2008),
    (3, "Record years before commissioning are ignored", 4_000_000.0, BEDLOAD, 100.0, 2003, 2006),
    (4, "Analysis year equals commissioning year", 4_000_000.0, BEDLOAD, 100.0, 2006, 2006),
    (5, "Dead volume already exceeded", 4_000_000.0, BEDLOAD, 100.0, 1990, 2006),
    (6, "Nothing delivered", 0.0, [0.0] * 5, 100.0, 2001, 2006),
    (7, "The minimum year delivers nothing", 0.0, SPARSE, 100.0, 2001, 2006),
]

NOT_REACHED, EXCEEDED = "not reached", "exceeded"


def elevation(volume):
    """H(V), limited to the volume at the crest."""
    return HBED + (HMAX - HBED) * (min(volume, VCREST) / VMAX) ** (1.0 / M)


def worked(case):
    _, _, suspended, bedload, te, commissioning, analysis = case
    totals = [te / 100.0 * (suspended + vb) / M3_PER_HM3 for vb in bedload]
    rates = {"Mean": sum(totals) / len(totals), "Min": min(totals), "Max": max(totals)}
    used, in_record, cumulative, running = [], [], [], 0.0
    for year in range(commissioning, analysis):
        known = year in RECORD_YEARS
        used.append(totals[RECORD_YEARS.index(year)] if known else rates["Mean"])
        in_record.append(1 if known else 0)
        running += used[-1]
        cumulative.append(running)
    s0 = sum(used)
    exceeded = None
    if s0 >= VDEAD:
        exceeded = next(y for y, c in zip(range(commissioning, analysis), cumulative, strict=True) if c >= VDEAD)

    def lifespan(rate):
        if s0 >= VDEAD:
            return EXCEEDED
        return NOT_REACHED if rate == 0 else analysis + (VDEAD - s0) / rate

    return totals, rates, used, in_record, cumulative, s0, exceeded, {name: lifespan(rates[name]) for name in RATES}


def write_csv():
    with open(HERE / "integration_cases.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["Case", "Description", "Hbed_m", "Hmax_m", "Hcrest_m", "Vactive_hm3", "Vmax_hm3", "m"]
            + ["Suspended_m3_yr", "TE_pct", "Commissioning_year", "Analysis_year", "Vdead_hm3", "Vcrest_hm3", "Hmin_m"]
            + [f"Rate_{name}_hm3_yr" for name in RATES]
            + ["S0_hm3", "S0_Elevation_m", "Exceeded_Year"]
            + [f"Lifespan_{name}" for name in RATES]
        )
        for case in CASES:
            number, description, suspended, _, te, commissioning, analysis = case
            _, rates, _, _, _, s0, exceeded, lifespans = worked(case)
            writer.writerow(
                [number, description, HBED, HMAX, HCREST, VACTIVE, VMAX, M, suspended, te, commissioning, analysis]
                + [VDEAD, VCREST, HMIN, *(rates[name] for name in RATES), s0, elevation(s0)]
                + ["" if exceeded is None else exceeded, *(lifespans[name] for name in RATES)]
            )

    with open(HERE / "integration_record.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Case", "Year", "Bedload_m3_yr", "Total_hm3", "Elevation_m"])
        for case in CASES:
            totals = worked(case)[0]
            for year, bedload, total in zip(RECORD_YEARS, case[3], totals, strict=True):
                writer.writerow([case[0], year, bedload, total, elevation(total)])

    with open(HERE / "integration_accumulation.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Case", "Year", "In_Record", "Total_Used_hm3", "Cumulative_hm3"])
        for case in CASES:
            _, _, used, in_record, cumulative, _, _, _ = worked(case)
            for year, flag, value, total in zip(range(case[5], case[6]), in_record, used, cumulative, strict=True):
                writer.writerow([case[0], year, flag, value, total])

    with open(HERE / "integration_projection.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Case", "Rate", "K", "Year", "Cumulative_hm3", "Elevation_m"])
        for case in CASES:
            _, rates, _, _, _, s0, _, _ = worked(case)
            for name in RATES:
                for k in HORIZONS:
                    volume = s0 + k * rates[name]
                    writer.writerow([case[0], name, k, case[6] + k, volume, elevation(volume)])


def case_sheet(book, case):
    number, description, suspended, bedload, te, commissioning, analysis = case
    sheet = book.create_sheet(f"Case {number}")
    sheet["A1"] = f"Case {number}: {description}"
    sheet["A1"].font = Font(bold=True)

    inputs = [
        ("Hbed (m)", HBED),
        ("Hmax (m)", HMAX),
        ("Hcrest (m)", HCREST),
        ("Vactive (hm3)", VACTIVE),
        ("Vmax (hm3)", VMAX),
        ("Exponent m", M),
        ("Suspended Vs (m3/yr)", suspended),
        ("Trap efficiency TE (%)", te),
        ("Commissioning year c", commissioning),
        ("Analysis year a", analysis),
    ]
    for row, (label, value) in enumerate(inputs, start=3):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
    derived = [
        ("Vdead (hm3) = Vmax - Vactive", "=B7-B6"),
        ("V(Hcrest) (hm3) = Vmax ((Hcrest-Hbed)/(Hmax-Hbed))^m", "=B7*((B5-B3)/(B4-B3))^B8"),
        ("Hmin (m) = H(Vdead)", "=B3+(B4-B3)*(B14/B7)^(1/B8)"),
    ]
    for row, (label, formula) in enumerate(derived, start=14):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=formula)

    for column, text in enumerate(
        ["Record year", "Bedload Vb (m3/yr)", "T (hm3) = TE/100 (Vs + Vb) / 1e6", "H(T) (m), limited to the crest"],
        start=1,
    ):
        sheet.cell(row=18, column=column, value=text).font = Font(bold=True)
    first, last = 19, 18 + len(RECORD_YEARS)
    for row, (year, vb) in enumerate(zip(RECORD_YEARS, bedload, strict=True), start=first):
        sheet.cell(row=row, column=1, value=year)
        sheet.cell(row=row, column=2, value=vb)
        sheet.cell(row=row, column=3, value=f"=$B$10/100*($B$9+B{row})/1000000")
        sheet.cell(row=row, column=4, value=f"=$B$3+($B$4-$B$3)*(MIN(C{row},$B$15)/$B$7)^(1/$B$8)")

    sheet["A25"], sheet["B25"] = "Rate, mean (hm3/yr)", f"=AVERAGE(C{first}:C{last})"
    sheet["A26"], sheet["B26"] = "Rate, min year (hm3/yr)", f"=MIN(C{first}:C{last})"
    sheet["A27"], sheet["B27"] = "Rate, max year (hm3/yr)", f"=MAX(C{first}:C{last})"

    count = analysis - commissioning
    acc_first, acc_last = 37, 36 + count
    sheet["A29"] = "Accumulation to date S0 (hm3) = sum of T' for c <= y < a"
    sheet["B29"] = f"=SUM(C{acc_first}:C{acc_last})" if count else 0
    sheet["A30"] = "H(S0) (m), limited to the crest"
    sheet["B30"] = "=$B$3+($B$4-$B$3)*(MIN(B29,$B$15)/$B$7)^(1/$B$8)"
    sheet["A31"] = "Year the dead volume was passed (if S0 >= Vdead)"
    sheet["B31"] = f'=IF(B29>=B14,MIN(E{acc_first}:E{acc_last}),"no")' if count else "no"
    for row, (name, source) in enumerate(zip(RATES, (25, 26, 27), strict=True), start=32):
        sheet.cell(row=row, column=1, value=f"Lifespan at the {name.lower()} rate: a + (Vdead - S0) / r")
        sheet.cell(
            row=row,
            column=2,
            value=f'=IF($B$29>=$B$14,"exceeded",IF(B{source}=0,"not reached",$B$12+($B$14-$B$29)/B{source}))',
        )

    for column, text in enumerate(
        ["Year c to a-1", "In record (1/0)", "T' used (hm3)", "Cumulative (hm3)", "Year, or +10000 if below Vdead"],
        start=1,
    ):
        sheet.cell(row=36, column=column, value=text).font = Font(bold=True)
    record_years, record_totals = f"$A${first}:$A${last}", f"$C${first}:$C${last}"
    for index, row in enumerate(range(acc_first, acc_last + 1)):
        sheet.cell(row=row, column=1, value=commissioning + index)
        sheet.cell(row=row, column=2, value=f"=COUNTIF({record_years},A{row})")
        sheet.cell(row=row, column=3, value=f"=SUMIF({record_years},A{row},{record_totals})+(1-B{row})*$B$25")
        sheet.cell(row=row, column=4, value=f"=C{row}" if index == 0 else f"=D{row - 1}+C{row}")
        sheet.cell(row=row, column=5, value=f"=A{row}+(1-(D{row}>=$B$14))*10000")

    top = max(acc_last, 36) + 3
    for column, text in enumerate(
        ["Rate", "Years after a (k)", "Year a + k", "r (hm3/yr)", "S0 + k r (hm3)", "H(S) (m), limited to the crest"],
        start=1,
    ):
        sheet.cell(row=top, column=column, value=text).font = Font(bold=True)
    row = top + 1
    for name, source in zip(RATES, (25, 26, 27), strict=True):
        for k in HORIZONS:
            sheet.cell(row=row, column=1, value=name)
            sheet.cell(row=row, column=2, value=k)
            sheet.cell(row=row, column=3, value=f"=$B$12+B{row}")
            sheet.cell(row=row, column=4, value=f"=$B${source}")
            sheet.cell(row=row, column=5, value=f"=$B$29+B{row}*D{row}")
            sheet.cell(row=row, column=6, value=f"=$B$3+($B$4-$B$3)*(MIN(E{row},$B$15)/$B$7)^(1/$B$8)")
            row += 1
    for column, width in zip("ABCDEF", (58, 20, 34, 30, 28, 30), strict=True):
        sheet.column_dimensions[column].width = width


def write_workbook():
    book = Workbook()
    book.remove(book.active)
    for case in CASES:
        case_sheet(book, case)
    book.calculation.fullCalcOnLoad = True  # openpyxl stores no results, so Excel works them out on opening
    book.save(HERE / "integration_verification.xlsx")


if __name__ == "__main__":
    write_csv()
    write_workbook()
