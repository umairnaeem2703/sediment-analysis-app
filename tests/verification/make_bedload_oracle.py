"""Write the M3 verification oracle for the Wilcock & Crowe bedload formulas (B3.2; OI-04).

Run `python tests/verification/make_bedload_oracle.py` to rewrite these files next to this script:

- `bedload_verification.xlsx`: the hand calculation as live spreadsheet formulas, for the hydrologist to open and
  review. The constants sit on their own sheet; every step of B3.2 is one column.
- `bedload_wstar.csv`, `bedload_tau_ri.csv`, `bedload_transport.csv`, `bedload_series.csv`, `bedload_annual.csv`:
  the same cases with the results worked out below using only `math`. The unit tests read these files, so the oracle
  never depends on `src`.

The grain fractions are the sample file's (inputs/fractional_parameters.csv). Change a case or a constant here, review
the workbook, and only then commit the files.
"""

import csv
import math
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

G, RHO_W, RHO_S = 9.81, 1000.0, 2650.0
H, B, N = 0.5, 50.0, 0.05  # flow depth (m), bed width (m), Manning n
D50, FS = 0.02, 0.10
PHI_BREAK = 1.35
SECONDS_PER_YEAR = 31_557_600.0  # 365.25 days
S = RHO_S / RHO_W

FRACTIONS = [(0.004, 0.10), (0.010, 0.20), (0.020, 0.35), (0.032, 0.25), (0.050, 0.10)]  # (di in m, Fi)
EXTRA_DIAMETERS = [0.001, 0.100]  # tau_ri checks beyond the sample fractions

PHI_CASES = [0.0, 0.1, 0.5, 1.0, 1.3, 1.35, 1.5, 2.0, 4.0, 10.0]
FLOW_CASES = [0.0, 5.0, 10.0, 12.5, 15.0, 25.0, 50.0, 100.0]  # m3/s
SERIES = [  # (date, flow m3/s): two calendar years, three water years
    (date(2009, 9, 30), 12.0),
    (date(2009, 10, 1), 18.5),
    (date(2009, 12, 31), 0.0),
    (date(2010, 1, 1), 9.2),
    (date(2010, 6, 15), 25.0),
    (date(2010, 9, 30), 14.4),
    (date(2010, 10, 1), 31.0),
    (date(2010, 12, 31), 6.5),
]

TAU_RM_STAR = 0.021 + 0.015 * math.exp(-20 * FS)


def wstar(phi):
    if phi <= 0:
        return 0.0
    if phi < PHI_BREAK:
        return 0.002 * phi**7.5
    return 14.0 * (1.0 - 0.894 / math.sqrt(phi)) ** 4.5


def exponent(di):
    return 0.67 / (1.0 + math.exp(1.5 - di / D50))


def tau_ri(di):
    """The hydrologist's form (D50): the hiding function scales the Shields number, converted to Pa with di."""
    return TAU_RM_STAR * (di / D50) ** exponent(di) * (S - 1) * RHO_W * G * di


def hydraulics(q):
    velocity = q / (B * H)
    tau = RHO_W * G * N**2 * velocity**2 / H ** (1.0 / 3.0)
    return velocity, tau, math.sqrt(tau / RHO_W)


def transport(q):
    """Unit-width bedload (m2/s) summed over the sample fractions."""
    _, tau, u_star = hydraulics(q)
    return sum(wstar(tau / tau_ri(di)) * fi * u_star**3 / ((S - 1.0) * G) for di, fi in FRACTIONS)


def water_year(day):
    return day.year + 1 if day.month >= 10 else day.year


def annual_rows():
    """(grouping, key, mean flow, mean transport, volume m3/yr, mass t/yr) for each calendar and water year."""
    rows = []
    for grouping, key_of in (("calendar", lambda d: d.year), ("water", water_year)):
        for key in sorted({key_of(day) for day, _ in SERIES}):
            flows = [q for day, q in SERIES if key_of(day) == key]
            mean_flow = sum(flows) / len(flows)
            mean_qb = sum(transport(q) for q in flows) / len(flows)
            volume = mean_qb * B * SECONDS_PER_YEAR
            rows.append((grouping, key, mean_flow, mean_qb, volume, volume * RHO_S / 1000.0))
    return rows


def write_csvs():
    def write(name, header, rows):
        with open(HERE / name, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(header)
            writer.writerows(rows)

    write("bedload_wstar.csv", ["Phi", "W_star"], [(phi, wstar(phi)) for phi in PHI_CASES])
    write(
        "bedload_tau_ri.csv",
        ["di_m", "Fi", "b_i", "tau_ri_Pa"],
        [(di, fi, exponent(di), tau_ri(di)) for di, fi in FRACTIONS]
        + [(di, "", exponent(di), tau_ri(di)) for di in EXTRA_DIAMETERS],
    )
    transport_rows = []
    for q in FLOW_CASES:
        _, tau, u_star = hydraulics(q)
        transport_rows.append((q, tau, u_star, transport(q)))
    write("bedload_transport.csv", ["Flow_m3_s", "Tau_Pa", "U_star_m_s", "Qb_m2_s"], transport_rows)
    write(
        "bedload_series.csv",
        ["Date", "Flow", "Water_Year", "Qb_m2_s"],
        [(day.strftime("%d/%m/%Y"), q, water_year(day), transport(q)) for day, q in SERIES],
    )
    write(
        "bedload_annual.csv",
        ["Grouping", "Key", "Mean_Flow_m3_s", "Mean_Qb_m2_s", "Volume_m3_yr", "Mass_t_yr"],
        annual_rows(),
    )


def _bold(sheet, headers):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def write_workbook():
    book = Workbook()

    constants = book.active
    constants.title = "Constants"
    _bold(constants, ["Constant", "Value"])
    rows = [
        ("g (m/s2)", G),
        ("rho_w (kg/m3)", RHO_W),
        ("rho_s (kg/m3)", RHO_S),
        ("h, flow depth (m)", H),
        ("B, bed width (m)", B),
        ("n, Manning", N),
        ("D50 (m)", D50),
        ("Fs, sand fraction", FS),
        ("phi break", PHI_BREAK),
        ("Seconds per year (365.25 d)", SECONDS_PER_YEAR),
    ]
    for label, value in rows:
        constants.append([label, value])
    constants.append(["s = rho_s / rho_w", "=B4/B3"])
    constants.append(["tau_rm* = 0.021 + 0.015 exp(-20 Fs)", "=0.021+0.015*EXP(-20*B9)"])
    constants.append(["Mass factor rho_s / 1000 (t/m3)", "=B4/1000"])
    constants.column_dimensions["A"].width = 42
    g, rho_w, h, bed, n, d50, phi_break, per_year, s, tau_rm_star, mass = (
        f"Constants!$B${row}" for row in (2, 3, 5, 6, 7, 8, 10, 11, 12, 13, 14)
    )

    wstar_sheet = book.create_sheet("W star")
    _bold(
        wstar_sheet, ["phi", "W* (published)", "Low branch 0.002 phi^7.5", "High branch 14 (1 - 0.894/sqrt(phi))^4.5"]
    )
    for row, phi in enumerate(PHI_CASES, start=2):
        wstar_sheet.append(
            [
                phi,
                f"=IF(A{row}<=0,0,IF(A{row}<{phi_break},C{row},D{row}))",
                f"=0.002*A{row}^7.5",
                f"=IF(A{row}>0,14*(1-0.894/SQRT(A{row}))^4.5,0)",
            ]
        )
    for column, width in zip("ABCD", (8, 16, 26, 40), strict=True):
        wstar_sheet.column_dimensions[column].width = width

    fractions = book.create_sheet("Fractions")
    _bold(
        fractions,
        [
            "di (m)",
            "Fi",
            "di / D50",
            "b_i = 0.67 / (1 + exp(1.5 - di/D50))",
            "tau_ri* = tau_rm* (di/D50)^b_i",
            "tau_ri = tau_ri* (s-1) rho_w g di (Pa)",
        ],
    )
    for row, (di, fi) in enumerate(FRACTIONS + [(d, None) for d in EXTRA_DIAMETERS], start=2):
        fractions.append(
            [
                di,
                fi,
                f"=A{row}/{d50}",
                f"=0.67/(1+EXP(1.5-C{row}))",
                f"={tau_rm_star}*C{row}^D{row}",
                f"=E{row}*({s}-1)*{rho_w}*{g}*A{row}",
            ]
        )
    fractions["H1"] = "Rows without Fi only check tau_ri."
    for column, width in zip("ABCDEF", (8, 8, 10, 38, 30, 38), strict=True):
        fractions.column_dimensions[column].width = width
    first = 2  # the sheet row of the first sample fraction

    transport_sheet = book.create_sheet("Transport")
    _bold(
        transport_sheet,
        [
            "Q (m3/s)",
            "di (m)",
            "Fi",
            "v = Q / (B h)",
            "tau = rho_w g n^2 v^2 / h^(1/3) (Pa)",
            "u* = sqrt(tau / rho_w)",
            "tau_ri (Pa)",
            "phi = tau / tau_ri",
            "W*",
            "q_bi = W* Fi u*^3 / ((s-1) g)",
        ],
    )
    row = 2
    for q in FLOW_CASES:
        for k in range(len(FRACTIONS)):
            f = first + k
            transport_sheet.append(
                [
                    q,
                    f"=Fractions!$A${f}",
                    f"=Fractions!$B${f}",
                    f"=A{row}/({bed}*{h})",
                    f"={rho_w}*{g}*{n}^2*D{row}^2/{h}^(1/3)",
                    f"=SQRT(E{row}/{rho_w})",
                    f"=Fractions!$F${f}",
                    f"=E{row}/G{row}",
                    f"=IF(H{row}<=0,0,IF(H{row}<{phi_break},0.002*H{row}^7.5,14*(1-0.894/SQRT(H{row}))^4.5))",
                    f"=I{row}*C{row}*F{row}^3/(({s}-1)*{g})",
                ]
            )
            row += 1
    rows_end = row - 1
    transport_sheet["L1"], transport_sheet["M1"] = "Q (m3/s)", "q_b = sum of q_bi (m2/s)"
    for cell in (transport_sheet["L1"], transport_sheet["M1"]):
        cell.font = Font(bold=True)
    for offset, q in enumerate(FLOW_CASES, start=2):
        transport_sheet[f"L{offset}"] = q
        transport_sheet[f"M{offset}"] = f"=SUMIF($A$2:$A${rows_end},L{offset},$J$2:$J${rows_end})"
    for column, width in zip("ABCDEFGHIJLM", (9, 8, 6, 14, 34, 22, 12, 16, 8, 28, 10, 26), strict=True):
        transport_sheet.column_dimensions[column].width = width

    series = book.create_sheet("Annual")
    _bold(
        series,
        [
            "Date",
            "Q (m3/s)",
            "Water_Year",
            "v",
            "tau (Pa)",
            "u* (m/s)",
            *(f"q_b fraction {k + 1} (m2/s)" for k in range(len(FRACTIONS))),
            "q_b daily (m2/s)",
            "Calendar year",
        ],
    )
    for row, (day, q) in enumerate(SERIES, start=2):
        per_fraction = []
        for k in range(len(FRACTIONS)):
            f = first + k
            phi = f"E{row}/Fractions!$F${f}"
            per_fraction.append(
                f"=IF(E{row}<=0,0,IF({phi}<{phi_break},0.002*({phi})^7.5,14*(1-0.894/SQRT({phi}))^4.5))"
                f"*Fractions!$B${f}*F{row}^3/(({s}-1)*{g})"
            )
        series.append(
            [
                day,
                q,
                f"=IF(MONTH(A{row})>=10,YEAR(A{row})+1,YEAR(A{row}))",
                f"=B{row}/({bed}*{h})",
                f"={rho_w}*{g}*{n}^2*D{row}^2/{h}^(1/3)",
                f"=SQRT(E{row}/{rho_w})",
                *per_fraction,
                f"=SUM(G{row}:K{row})",
                f"=YEAR(A{row})",
            ]
        )
        series[f"A{row}"].number_format = "dd/mm/yyyy"
    end = 1 + len(SERIES)
    summary_headers = [
        "Grouping",
        "Key",
        "Mean Q (m3/s)",
        "Mean q_b (m2/s)",
        "Volume V_b = q_b B s/yr (m3/yr)",
        "Mass M_b (t/yr)",
    ]
    for offset, header in enumerate(summary_headers):
        cell = series.cell(row=1, column=15 + offset, value=header)
        cell.font = Font(bold=True)
    for row, (grouping, key, *_rest) in enumerate(annual_rows(), start=2):
        key_range = f"$C$2:$C${end}" if grouping == "water" else f"$M$2:$M${end}"
        values = [
            grouping,
            key,
            f"=AVERAGEIF({key_range},P{row},$B$2:$B${end})",
            f"=AVERAGEIF({key_range},P{row},$L$2:$L${end})",
            f"=R{row}*{bed}*{per_year}",
            f"=S{row}*{mass}",
        ]
        for offset, value in enumerate(values):
            series.cell(row=row, column=15 + offset, value=value)
    series.column_dimensions["A"].width = 12
    for column in "OPQRST":
        series.column_dimensions[column].width = 18
    series.column_dimensions["S"].width = 32

    book.calculation.fullCalcOnLoad = True  # openpyxl stores no results, so Excel works them out on opening
    book.save(HERE / "bedload_verification.xlsx")


if __name__ == "__main__":
    write_csvs()
    write_workbook()
