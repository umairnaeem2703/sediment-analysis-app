"""Write the M2 verification oracle for the suspended-sediment formulas (B3.1; OI-06).

Run `python tests/verification/make_suspended_oracle.py` to rewrite two files next to this script:

- `suspended_verification.xlsx`: the hand calculation as live spreadsheet formulas, for the hydrologist to open and
  review. The constants sit in their own cells; every step of B3.1 is one column.
- `suspended_cases.csv`: the same cases with the results worked out below using only `math`. The unit tests read
  this file, so the oracle never depends on `src`.

Change a case or a constant here, review the workbook, and only then commit both files.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

USDA_COEFFICIENT, USDA_EXPONENT = 0.51, -0.11
CEM_INTERCEPT, CEM_SLOPE_TERM = 2.05823, 0.02816
BULK_DENSITY_T_M3 = 1.3
HA_PER_KM2 = 100.0

# (case, description, zonal mean t/ha/yr, area km2, slope m/m)
CASES = [
    (1, "Kavsakkaya sample sub-basin", 3.148701574, 309.2, 0.23),
    (2, "Cubuk 1 sample sub-basin", 3.148701574, 834.7, 0.09),
    (3, "Cubuk 2 sample sub-basin", 3.148701574, 180.2, 0.16),
    (4, "Unit area, flat, round mean", 10.0, 1.0, 0.0),
    (5, "Mid-size steep basin", 2.5, 100.0, 0.5),
    (6, "Small very steep basin", 40.0, 0.25, 1.2),
    (7, "Zero erosion mean", 0.0, 50.0, 0.1),
]

RESULT_COLUMNS = [
    "Area_ha",
    "Gross_Mass_t",
    "SIO_USDA_SCS",
    "SIO_CEM",
    "SIO_Average",
    "Yield_Mass_t",
    "Volume_m3",
    "Specific_Yield_m3_km2",
]


def worked(mean, area, slope):
    area_ha = HA_PER_KM2 * area
    gross = mean * area_ha
    usda = USDA_COEFFICIENT * area**USDA_EXPONENT
    cem = 1.0 / (CEM_INTERCEPT + CEM_SLOPE_TERM * slope**2)
    average = (usda + cem) / 2.0
    yield_mass = gross * average
    volume = yield_mass / BULK_DENSITY_T_M3
    return [area_ha, gross, usda, cem, average, yield_mass, volume, volume / area]


def write_csv():
    header = ["Case", "Description", "Zonal_Mean_t_ha_yr", "Area_km2", "Slope_m_per_m", *RESULT_COLUMNS]
    with open(HERE / "suspended_cases.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for case, description, mean, area, slope in CASES:
            writer.writerow([case, description, mean, area, slope, *worked(mean, area, slope)])


def write_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "B3.1 cases"
    headers = [
        "Case",
        "Description",
        "Zonal mean E (t/ha/yr)",
        "Area A (km2)",
        "Slope S (m/m)",
        "Area (ha) = 100 A",
        "Gross erosion G (t/yr) = E x ha",
        "SIO USDA-SCS = 0.51 A^-0.11",
        "SIO CEM = 1 / (2.05823 + 0.02816 S^2)",
        "SIO average = (USDA + CEM) / 2",
        "Yield mass Y (t/yr) = G x SIO",
        "Volume Vs (m3/yr) = Y / 1.3",
        "Specific yield (m3/yr/km2) = Vs / A",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    ha, coefficient, exponent, intercept, slope_term, density = (f"$P${row}" for row in range(2, 8))

    for row, (case, description, mean, area, slope) in enumerate(CASES, start=2):
        sheet.append(
            [
                case,
                description,
                mean,
                area,
                slope,
                f"={ha}*D{row}",
                f"=C{row}*F{row}",
                f"={coefficient}*D{row}^{exponent}",
                f"=1/({intercept}+{slope_term}*E{row}^2)",
                f"=(H{row}+I{row})/2",
                f"=G{row}*J{row}",
                f"=K{row}/{density}",
                f"=L{row}/D{row}",
            ]
        )
    constants = [
        ("Hectares per km2", HA_PER_KM2),
        ("USDA-SCS coefficient", USDA_COEFFICIENT),
        ("USDA-SCS exponent", USDA_EXPONENT),
        ("CEM intercept", CEM_INTERCEPT),
        ("CEM slope term", CEM_SLOPE_TERM),
        ("Bulk density (t/m3)", BULK_DENSITY_T_M3),
    ]
    sheet["O1"], sheet["P1"] = "Constant", "Value"
    for row, (label, value) in enumerate(constants, start=2):
        sheet.cell(row=row, column=15, value=label)
        sheet.cell(row=row, column=16, value=value)
    for column, width in zip("ABCDEFGHIJKLM", (6, 32, 14, 12, 12, 16, 22, 22, 28, 22, 22, 22, 24), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["O"].width = 24
    book.calculation.fullCalcOnLoad = True  # openpyxl stores no results, so Excel works them out on opening
    book.save(HERE / "suspended_verification.xlsx")


if __name__ == "__main__":
    write_csv()
    write_workbook()
